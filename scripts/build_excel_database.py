"""
Complete Master Excel Builder and Quality Auditor for StudentTools.cyou
Generates 1,105 High-Quality Topics across 5 Tabs with Formatting, AutoFilters, and Data Validation.
"""

import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from topic_generator.gen_jee import get_jee_topics
from topic_generator.gen_student import get_student_topics
from topic_generator.gen_career import get_career_topics
from topic_generator.gen_finance import get_finance_topics
from topic_generator.gen_general import get_general_topics

# Canonical 38 Tools mapping from data/tools.ts
TOOLS_CATALOGUE = [
    # JEE (6)
    {"name": "JEE Main Marks to Percentile", "slug": "/jee/marks-to-percentile", "category": "JEE"},
    {"name": "JEE Main Percentile to Rank", "slug": "/jee/percentile-to-rank", "category": "JEE"},
    {"name": "JEE College Predictor", "slug": "/jee/college-predictor", "category": "JEE"},
    {"name": "JEE Cutoff Explorer", "slug": "/jee/cutoff", "category": "JEE"},
    {"name": "JEE Marks Calculator", "slug": "/jee/marks-calculator", "category": "JEE"},
    {"name": "JEE Exam Countdown", "slug": "/jee/countdown", "category": "JEE"},
    # Student (8)
    {"name": "CGPA to Percentage Calculator", "slug": "/student/cgpa-to-percentage", "category": "Student"},
    {"name": "Percentage Calculator", "slug": "/student/percentage-calculator", "category": "Student"},
    {"name": "Marks Percentage Calculator", "slug": "/student/marks-percentage", "category": "Student"},
    {"name": "GPA Calculator", "slug": "/student/gpa-calculator", "category": "Student"},
    {"name": "Attendance Calculator", "slug": "/student/attendance-calculator", "category": "Student"},
    {"name": "Required Marks Calculator", "slug": "/student/required-marks", "category": "Student"},
    {"name": "Average Marks Calculator", "slug": "/student/average-marks", "category": "Student"},
    {"name": "Study Hours Calculator", "slug": "/student/study-hours", "category": "Student"},
    # Career (8)
    {"name": "CTC to In-Hand Salary Calculator", "slug": "/career/ctc-to-in-hand", "category": "Career"},
    {"name": "Monthly to Annual Salary Calculator", "slug": "/career/monthly-to-annual-salary", "category": "Career"},
    {"name": "Annual to Monthly Salary Calculator", "slug": "/career/annual-to-monthly-salary", "category": "Career"},
    {"name": "Salary Hike Calculator", "slug": "/career/salary-hike", "category": "Career"},
    {"name": "Increment Calculator", "slug": "/career/increment-calculator", "category": "Career"},
    {"name": "PF Calculator", "slug": "/career/pf-calculator", "category": "Career"},
    {"name": "Bonus Calculator", "slug": "/career/bonus-calculator", "category": "Career"},
    {"name": "Internship Stipend Calculator", "slug": "/career/internship-stipend", "category": "Career"},
    # Finance (6)
    {"name": "EMI Calculator", "slug": "/finance/emi-calculator", "category": "Finance"},
    {"name": "SIP Calculator", "slug": "/finance/sip-calculator", "category": "Finance"},
    {"name": "GST Calculator", "slug": "/finance/gst-calculator", "category": "Finance"},
    {"name": "FD Calculator", "slug": "/finance/fd-calculator", "category": "Finance"},
    {"name": "RD Calculator", "slug": "/finance/rd-calculator", "category": "Finance"},
    {"name": "Income Tax Calculator", "slug": "/finance/tax-calculator", "category": "Finance"},
    # General (10)
    {"name": "Age Calculator", "slug": "/calculators/age", "category": "General"},
    {"name": "BMI Calculator", "slug": "/calculators/bmi", "category": "General"},
    {"name": "Scientific Calculator", "slug": "/calculators/scientific-calculator", "category": "General"},
    {"name": "Unit Converter", "slug": "/calculators/unit-converter", "category": "General"},
    {"name": "Date Calculator", "slug": "/calculators/date", "category": "General"},
    {"name": "Discount Calculator", "slug": "/calculators/discount", "category": "General"},
    {"name": "Profit and Loss Calculator", "slug": "/calculators/profit-loss", "category": "General"},
    {"name": "Ratio Calculator", "slug": "/calculators/ratio", "category": "General"},
    {"name": "Fraction Calculator", "slug": "/calculators/fraction", "category": "General"},
    {"name": "Percentage Change Calculator", "slug": "/calculators/percentage-change", "category": "General"},
]

def build_all_topics():
    print("Collecting topics from modules...")
    jee = get_jee_topics()
    student = get_student_topics()
    career = get_career_topics()
    finance = get_finance_topics()
    general = get_general_topics()
    
    all_raw = jee + student + career + finance + general
    print(f"Total raw topics collected: {len(all_raw)}")
    print(f" - JEE: {len(jee)}")
    print(f" - Student: {len(student)}")
    print(f" - Career: {len(career)}")
    print(f" - Finance: {len(finance)}")
    print(f" - General: {len(general)}")
    
    # Audit & Assign IDs
    seen_topics = set()
    final_topics = []
    
    for idx, item in enumerate(all_raw, start=1):
        topic_id = f"ST{idx:04d}" if idx >= 1000 else f"ST{idx:03d}"
        title = item["topic"].strip()
        title_lower = title.lower()
        
        if title_lower in seen_topics:
            raise ValueError(f"Duplicate topic title detected: {title}")
        seen_topics.add(title_lower)
        
        # Validation checks
        assert item["category"] in ["JEE", "Student", "Career", "Finance", "General"], f"Invalid category: {item['category']}"
        assert item["article_type"] in ["Guide", "How-To", "Explanation", "Formula Guide", "Calculator Guide", "Example", "Comparison", "FAQ", "Planning Guide", "Reference"], f"Invalid article_type: {item['article_type']}"
        assert item["search_intent"] in ["Informational", "How-To", "Calculator", "Educational", "Comparison", "Planning", "Example", "Reference"], f"Invalid search_intent: {item['search_intent']}"
        assert item["evergreen"] in ["YES", "NO", "SEASONAL"], f"Invalid evergreen: {item['evergreen']}"
        assert item["priority"] in ["HIGH", "MEDIUM", "LOW"], f"Invalid priority: {item['priority']}"
        
        # Verify tool slug exists in catalogue
        tool_slug = item["related_tool_slug"]
        matching_tool = [t for t in TOOLS_CATALOGUE if t["slug"] == tool_slug]
        assert len(matching_tool) > 0, f"Unknown tool slug: {tool_slug} in topic {title}"
        
        row = {
            "topic_id": topic_id,
            "topic": title,
            "category": item["category"],
            "topic_cluster": item["topic_cluster"],
            "article_type": item["article_type"],
            "search_intent": item["search_intent"],
            "primary_keyword": item["primary_keyword"].strip(),
            "secondary_keywords": item["secondary_keywords"].strip(),
            "related_tool": item["related_tool"].strip(),
            "related_tool_slug": tool_slug,
            "content_angle": item["content_angle"].strip(),
            "evergreen": item["evergreen"],
            "priority": item["priority"],
            "status": "READY",
            "published_url": "",
            "published_at": "",
            "error": "",
            "content_hash": ""
        }
        final_topics.append(row)
        
    print(f"Audit passed cleanly! {len(final_topics)} unique, validated topics ready.")
    return final_topics

def create_excel_file(topics, output_path):
    print(f"Creating Excel workbook at {output_path}...")
    wb = openpyxl.Workbook()
    
    # Setup Styles
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Deep Navy
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=10)
    data_font_bold = Font(name="Segoe UI", size=10, bold=True)
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid") # Slate-50
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    
    thin_border_side = Side(border_style="thin", color="E2E8F0")
    thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    
    # -------------------------------------------------------------
    # TAB 1: Topics
    # -------------------------------------------------------------
    ws_topics = wb.active
    ws_topics.title = "Topics"
    ws_topics.views.sheetView[0].showGridLines = True
    
    headers = [
        "topic_id", "topic", "category", "topic_cluster", "article_type",
        "search_intent", "primary_keyword", "secondary_keywords", "related_tool",
        "related_tool_slug", "content_angle", "evergreen", "priority",
        "status", "published_url", "published_at", "error", "content_hash"
    ]
    
    ws_topics.append(headers)
    
    for col_num in range(1, len(headers) + 1):
        cell = ws_topics.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
        cell.border = thin_border
        
    ws_topics.row_dimensions[1].height = 28
    
    import re
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

    for row_idx, t in enumerate(topics, start=2):
        row_vals = [
            t["topic_id"], t["topic"], t["category"], t["topic_cluster"], t["article_type"],
            t["search_intent"], t["primary_keyword"], t["secondary_keywords"], t["related_tool"],
            t["related_tool_slug"], t["content_angle"], t["evergreen"], t["priority"],
            t["status"], t["published_url"], t["published_at"], t["error"], t["content_hash"]
        ]
        # Clean illegal characters if string
        clean_row_vals = [
            ILLEGAL_CHARACTERS_RE.sub("", str(v)) if isinstance(v, str) else v
            for v in row_vals
        ]
        ws_topics.append(clean_row_vals)
        ws_topics.row_dimensions[row_idx].height = 20
        
        is_even = (row_idx % 2 == 0)
        curr_fill = white_fill if is_even else zebra_fill
        
        for col_idx in range(1, len(headers) + 1):
            cell = ws_topics.cell(row=row_idx, column=col_idx)
            cell.font = data_font
            cell.fill = curr_fill
            cell.border = thin_border
            
            # Alignments
            if col_idx in [1, 3, 5, 6, 12, 13, 14]: # ID, Category, Type, Intent, Evergreen, Priority, Status
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx in [2, 4, 7, 8, 9, 10, 11]:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
                
    # Freeze Header & Enable AutoFilter
    ws_topics.freeze_panes = "A2"
    ws_topics.auto_filter.ref = f"A1:R{len(topics) + 1}"
    
    # Dropdown Data Validations
    dv_cat = DataValidation(type="list", formula1='"JEE,Student,Career,Finance,General"', allow_blank=False)
    ws_topics.add_data_validation(dv_cat)
    dv_cat.add(f"C2:C{len(topics) + 1}")
    
    dv_type = DataValidation(type="list", formula1='"Guide,How-To,Explanation,Formula Guide,Calculator Guide,Example,Comparison,FAQ,Planning Guide,Reference"', allow_blank=False)
    ws_topics.add_data_validation(dv_type)
    dv_type.add(f"E2:E{len(topics) + 1}")
    
    dv_intent = DataValidation(type="list", formula1='"Informational,How-To,Calculator,Educational,Comparison,Planning,Example,Reference"', allow_blank=False)
    ws_topics.add_data_validation(dv_intent)
    dv_intent.add(f"F2:F{len(topics) + 1}")
    
    dv_evergreen = DataValidation(type="list", formula1='"YES,NO,SEASONAL"', allow_blank=False)
    ws_topics.add_data_validation(dv_evergreen)
    dv_evergreen.add(f"L2:L{len(topics) + 1}")
    
    dv_priority = DataValidation(type="list", formula1='"HIGH,MEDIUM,LOW"', allow_blank=False)
    ws_topics.add_data_validation(dv_priority)
    dv_priority.add(f"M2:M{len(topics) + 1}")
    
    dv_status = DataValidation(type="list", formula1='"READY,PROCESSING,PUBLISHED,FAILED,REVIEW_REQUIRED,SKIPPED"', allow_blank=False)
    ws_topics.add_data_validation(dv_status)
    dv_status.add(f"N2:N{len(topics) + 1}")

    # Set Column Widths
    col_widths = {
        "A": 12,  # topic_id
        "B": 50,  # topic
        "C": 14,  # category
        "D": 22,  # topic_cluster
        "E": 18,  # article_type
        "F": 16,  # search_intent
        "G": 32,  # primary_keyword
        "H": 36,  # secondary_keywords
        "I": 30,  # related_tool
        "J": 30,  # related_tool_slug
        "K": 65,  # content_angle
        "L": 14,  # evergreen
        "M": 14,  # priority
        "N": 16,  # status
        "O": 25,  # published_url
        "P": 20,  # published_at
        "Q": 20,  # error
        "R": 22   # content_hash
    }
    for col_letter, width in col_widths.items():
        ws_topics.column_dimensions[col_letter].width = width

    # -------------------------------------------------------------
    # TAB 2: Clusters
    # -------------------------------------------------------------
    ws_clusters = wb.create_sheet(title="Clusters")
    ws_clusters.views.sheetView[0].showGridLines = True
    
    cluster_headers = ["Category", "Topic Cluster", "Topic Count", "Key Themes & Search Direction"]
    ws_clusters.append(cluster_headers)
    ws_clusters.row_dimensions[1].height = 28
    
    for c in range(1, 5):
        cell = ws_clusters.cell(row=1, column=c)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
        
    cluster_summary = {}
    for t in topics:
        key = (t["category"], t["topic_cluster"])
        cluster_summary[key] = cluster_summary.get(key, 0) + 1
        
    cluster_descriptions = {
        ("JEE", "Marks & Scoring"): "Raw marks tallying, shift difficulty variance, negative marking analysis, score milestones (50-270).",
        ("JEE", "Percentile"): "NTA normalization formula, 7 decimals, session 1 vs 2, percentile bands (75-99.95), shift equity.",
        ("JEE", "Rank"): "AIR formula, CRL vs Category, tie-breakers, candidate volume impact, rank milestones (1k to 100k+).",
        ("JEE", "College Prediction"): "JoSAA counselling, HS vs OS quotas, branch comparison, closing ranks across 31 NITs, IIITs, and GFTIs.",
        ("JEE", "Exam Planning"): "Revision roadmaps, exam countdown schedules, sleep conditioning, stress triage, mock test strategy.",
        ("JEE", "Marks Calculator"): "Answer key accuracy, Section B numerical rules, opportunity cost math, error notebooks.",
        ("Student", "Attendance"): "75% criteria, recovery formulas, safe bunk calculations, medical condonation, university detention rules.",
        ("Student", "CGPA"): "University-specific multipliers (CBSE 9.5x, AICTE, VTU, MU, AU, KTU), milestones (6.0 to 10.0), transcript rules.",
        ("Student", "Percentage"): "Core percentage formulas, increase/decrease, reverse percentage, mental math benchmarks, consecutive percentages.",
        ("Student", "Marks"): "Class 10/12 board calculations, best 4/5 rules, internal vs practical splits, degree aggregate percentages.",
        ("Student", "GPA"): "Credit-weighted SGPA/CGPA, quality points, relative grading curves, backlog clearance impact, CBCS rules.",
        ("Student", "Required Marks"): "Target aggregate planning, end-sem recovery calculations, passing criteria with separate theory/practicals.",
        ("Student", "Average Marks"): "Arithmetic mean vs weighted average, running averages, outlier impacts, dropping lowest quiz scores.",
        ("Student", "Study Planning"): "Available study hours budget, subject credit allocation, 50/10 study blocks, dual-prep scheduling.",
        ("Career", "CTC"): "Gross vs net take-home, Indian salary slip line items (Basic, HRA, PF, PT, Gratuity, TDS), 3-35 LPA breakdowns.",
        ("Career", "Salary"): "Annual to monthly conversion, gross paycheck calculations, weekly/hourly rates, freelance projecting.",
        ("Career", "Salary Hike"): "Appraisal percentages, job switch negotiations, tax bracket progression, 5% to 100% hike models.",
        ("Career", "Increment"): "Absolute increments, backdated arrears, real wage growth vs inflation, CAGR over multi-year careers.",
        ("Career", "PF"): "EPFO 12% calculation, EPS pension formula, VPF, Rs 2.5L tax limit, UAN online transfer, retirement corpus math.",
        ("Career", "Bonus"): "Performance multipliers, heavy TDS deductions, statutory 8.33% bonus, pro-rata payouts, clawbacks.",
        ("Career", "Internship"): "Hourly/monthly stipend conversions, taxability under Section 10(16), living expense budgeting, PPO evaluation.",
        ("Finance", "EMI"): "Reducing balance math, amortization schedule, loan tenure vs interest trade-off, prepaying 1 extra EMI/yr.",
        ("Finance", "SIP"): "Mutual fund future value formula, rupee cost averaging, XIRR, Step-Up SIP 2x multiplier, Rs 1 Crore roadmap.",
        ("Finance", "GST"): "Exclusive vs inclusive formulas, reverse GST extraction from MRP, 50:50 CGST/SGST split, 5/12/18/28% slabs.",
        ("Finance", "FD"): "Quarterly compounding formula, cumulative vs non-cumulative, Section 194A TDS, DICGC Rs 5L safety net.",
        ("Finance", "RD"): "Series compounding on monthly deposits, Post Office 5-year RD, emergency fund accumulation, flexi RDs.",
        ("Finance", "Income Tax"): "New vs Old regime slabs FY 2024-25/2025-26, Section 87A rebate, standard deduction Rs 75k, 80C/80D/80CCD.",
        ("General", "Age"): "Day/month borrowing calendar math, government exam cutoffs, leap years, Zeller's congruence weekday finding.",
        ("General", "BMI"): "Metric/imperial formulas, Asian Indian lower cutoffs (23/25), waist-to-height ratio, muscle vs fat limitations.",
        ("General", "Unit Conversion"): "Length, temperature, mass, speed, digital binary storage (GB vs GiB), pressure, area (sq ft, acre, bigha).",
        ("General", "Date"): "Calendar intervals, business days excluding holidays, deadline addition, 90-day vs 3-month variance, Doomsday rule.",
        ("General", "Discount"): "Sale prices, successive chain discounts (50%+50%=75%), BOGO math, cash vs trade discounts, reverse markups.",
        ("General", "Profit/Loss"): "CP/SP formulas, markup on cost vs margin on revenue, break-even point in units, dishonest dealer math.",
        ("General", "Ratio"): "GCD simplification, extreme vs mean proportion, 16:9 aspect ratios, dividing amounts, golden ratio math.",
        ("General", "Fraction"): "Adding unlike denominators (LCD), multiplying, Keep-Change-Flip division, improper vs mixed conversions.",
    }
    
    row_num = 2
    for (cat, cluster), count in sorted(cluster_summary.items()):
        desc = cluster_descriptions.get((cat, cluster), "Practical calculations and search intents.")
        ws_clusters.append([cat, cluster, count, desc])
        ws_clusters.row_dimensions[row_num].height = 20
        
        is_even = (row_num % 2 == 0)
        curr_fill = white_fill if is_even else zebra_fill
        
        for c in range(1, 5):
            cell = ws_clusters.cell(row=row_num, column=c)
            cell.font = data_font
            cell.fill = curr_fill
            cell.border = thin_border
            if c in [1, 2]:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            elif c == 3:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.font = data_font_bold
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
        row_num += 1
        
    ws_clusters.column_dimensions["A"].width = 16
    ws_clusters.column_dimensions["B"].width = 24
    ws_clusters.column_dimensions["C"].width = 16
    ws_clusters.column_dimensions["D"].width = 90
    ws_clusters.freeze_panes = "A2"

    # -------------------------------------------------------------
    # TAB 3: Tools
    # -------------------------------------------------------------
    ws_tools = wb.create_sheet(title="Tools")
    ws_tools.views.sheetView[0].showGridLines = True
    
    tool_headers = ["Tool #", "Tool Name", "Category", "Canonical Route Slug", "Associated Topic Count"]
    ws_tools.append(tool_headers)
    ws_tools.row_dimensions[1].height = 28
    
    for c in range(1, 6):
        cell = ws_tools.cell(row=1, column=c)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
        
    tool_counts = {}
    for t in topics:
        tool_counts[t["related_tool_slug"]] = tool_counts.get(t["related_tool_slug"], 0) + 1
        
    for idx, tool in enumerate(TOOLS_CATALOGUE, start=1):
        cnt = tool_counts.get(tool["slug"], 0)
        ws_tools.append([idx, tool["name"], tool["category"], tool["slug"], cnt])
        curr_row = idx + 1
        ws_tools.row_dimensions[curr_row].height = 20
        
        is_even = (curr_row % 2 == 0)
        curr_fill = white_fill if is_even else zebra_fill
        
        for c in range(1, 6):
            cell = ws_tools.cell(row=curr_row, column=c)
            cell.font = data_font
            cell.fill = curr_fill
            cell.border = thin_border
            if c in [1, 3, 5]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
                
    ws_tools.column_dimensions["A"].width = 10
    ws_tools.column_dimensions["B"].width = 38
    ws_tools.column_dimensions["C"].width = 16
    ws_tools.column_dimensions["D"].width = 36
    ws_tools.column_dimensions["E"].width = 24
    ws_tools.freeze_panes = "A2"

    # -------------------------------------------------------------
    # TAB 4: Instructions
    # -------------------------------------------------------------
    ws_inst = wb.create_sheet(title="Instructions")
    ws_inst.views.sheetView[0].showGridLines = True
    
    ws_inst.column_dimensions["A"].width = 6
    ws_inst.column_dimensions["B"].width = 28
    ws_inst.column_dimensions["C"].width = 85
    
    inst_title_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid") # Slate-900
    inst_title_font = Font(name="Segoe UI", size=14, bold=True, color="FFFFFF")
    
    ws_inst.merge_cells("B1:C1")
    title_cell = ws_inst["B1"]
    title_cell.value = "StudentTools.cyou — Automated Content Publishing Pipeline Instructions"
    title_cell.fill = inst_title_fill
    title_cell.font = inst_title_font
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws_inst.row_dimensions[1].height = 36
    
    instructions_data = [
        ("Architecture Overview", "This spreadsheet serves as the master editorial queue for the StudentTools.cyou automated AI publishing engine."),
        ("Status Lifecycle", "The publication pipeline progresses strictly through: READY -> PROCESSING -> PUBLISHED (or FAILED / REVIEW_REQUIRED / SKIPPED)."),
        ("Worker Step 1: Select Topic", "Find the first row in the 'Topics' sheet where status == 'READY' and priority == 'HIGH' (or next available)."),
        ("Worker Step 2: Lock Record", "Immediately update the topic row status to 'PROCESSING' to prevent race conditions during concurrent runs."),
        ("Worker Step 3: Prompt Gemini", "Pass the topic title, article_type, search_intent, content_angle, and Indian context guidelines to Gemini."),
        ("Worker Step 4: Tool Injection", "Ensure the article embeds a contextual CTA and link to related_tool using the exact canonical related_tool_slug."),
        ("Worker Step 5: Quality Audit", "Perform automated linting: verify minimum word count (1,200+), proper Markdown formatting, zero AI hallucinatory figures, and accurate formulas."),
        ("Worker Step 6: Publish Post", "Deploy the article to the Next.js CMS/database, generating the live URL and canonical route."),
        ("Worker Step 7: Finalize Row", "Update status to 'PUBLISHED', write the live URL to published_url, set published_at timestamp, and record SHA-256 in content_hash."),
        ("Worker Error Handling", "If generation or linting fails, set status to 'FAILED', record the error trace in error, and continue to the next READY topic."),
        ("Editorial Independence", "Every row is fully distinct and self-contained; Gemini receives exact content angles to prevent duplicate or thin keyword articles."),
        ("Canonical Slugs", "Never alter related_tool_slug outside the 38 tools defined in the 'Tools' tab to maintain pristine internal linking integrity."),
    ]
    
    for i, (title, desc) in enumerate(instructions_data, start=3):
        ws_inst.cell(row=i, column=2, value=title).font = Font(name="Segoe UI", size=11, bold=True, color="1E3A8A")
        ws_inst.cell(row=i, column=3, value=desc).font = Font(name="Segoe UI", size=10)
        ws_inst.cell(row=i, column=2).border = thin_border
        ws_inst.cell(row=i, column=3).border = thin_border
        ws_inst.cell(row=i, column=2).alignment = Alignment(vertical="center")
        ws_inst.cell(row=i, column=3).alignment = Alignment(vertical="center", wrap_text=True)
        ws_inst.row_dimensions[i].height = 26

    # -------------------------------------------------------------
    # TAB 5: Stats
    # -------------------------------------------------------------
    ws_stats = wb.create_sheet(title="Stats")
    ws_stats.views.sheetView[0].showGridLines = True
    
    ws_stats.column_dimensions["A"].width = 5
    ws_stats.column_dimensions["B"].width = 28
    ws_stats.column_dimensions["C"].width = 18
    ws_stats.column_dimensions["D"].width = 5
    ws_stats.column_dimensions["E"].width = 28
    ws_stats.column_dimensions["F"].width = 18
    
    # Section 1: System Status
    ws_stats.merge_cells("B1:C1")
    s1 = ws_stats["B1"]
    s1.value = "Pipeline Status Summary"
    s1.fill = header_fill
    s1.font = header_font
    s1.alignment = Alignment(horizontal="center", vertical="center")
    ws_stats.row_dimensions[1].height = 28
    
    status_metrics = [
        ("Total Topics in Database", f"=COUNTA(Topics!A2:A{len(topics)+1})"),
        ("READY for Publishing", f'=COUNTIF(Topics!N2:N{len(topics)+1}, "READY")'),
        ("PROCESSING (In Progress)", f'=COUNTIF(Topics!N2:N{len(topics)+1}, "PROCESSING")'),
        ("PUBLISHED (Live)", f'=COUNTIF(Topics!N2:N{len(topics)+1}, "PUBLISHED")'),
        ("FAILED (Needs Retry)", f'=COUNTIF(Topics!N2:N{len(topics)+1}, "FAILED")'),
        ("REVIEW_REQUIRED", f'=COUNTIF(Topics!N2:N{len(topics)+1}, "REVIEW_REQUIRED")'),
        ("SKIPPED", f'=COUNTIF(Topics!N2:N{len(topics)+1}, "SKIPPED")'),
    ]
    
    for r_idx, (label, formula) in enumerate(status_metrics, start=2):
        ws_stats.cell(row=r_idx, column=2, value=label).font = data_font_bold if r_idx == 2 else data_font
        c_val = ws_stats.cell(row=r_idx, column=3, value=formula)
        c_val.font = data_font_bold
        c_val.alignment = Alignment(horizontal="center", vertical="center")
        ws_stats.cell(row=r_idx, column=2).border = thin_border
        c_val.border = thin_border
        ws_stats.row_dimensions[r_idx].height = 20

    # Section 2: Topics per Category
    ws_stats.merge_cells("E1:F1")
    s2 = ws_stats["E1"]
    s2.value = "Topics per Category"
    s2.fill = header_fill
    s2.font = header_font
    s2.alignment = Alignment(horizontal="center", vertical="center")
    
    cat_metrics = [
        ("JEE & Competitive Exams", '=COUNTIF(Topics!C2:C' + str(len(topics)+1) + ', "JEE")'),
        ("Student & Academic Tools", '=COUNTIF(Topics!C2:C' + str(len(topics)+1) + ', "Student")'),
        ("Career & Salary Tools", '=COUNTIF(Topics!C2:C' + str(len(topics)+1) + ', "Career")'),
        ("Finance Calculators", '=COUNTIF(Topics!C2:C' + str(len(topics)+1) + ', "Finance")'),
        ("General Calculators", '=COUNTIF(Topics!C2:C' + str(len(topics)+1) + ', "General")'),
    ]
    
    for r_idx, (label, formula) in enumerate(cat_metrics, start=2):
        ws_stats.cell(row=r_idx, column=5, value=label).font = data_font
        c_val = ws_stats.cell(row=r_idx, column=6, value=formula)
        c_val.font = data_font_bold
        c_val.alignment = Alignment(horizontal="center", vertical="center")
        ws_stats.cell(row=r_idx, column=5).border = thin_border
        c_val.border = thin_border
        ws_stats.row_dimensions[r_idx].height = 20
        
    # Save Workbook
    wb.save(output_path)
    print(f"Workbook successfully saved to {output_path}!")

if __name__ == "__main__":
    topics = build_all_topics()
    output_xlsx = os.path.abspath("StudentTools_1105_Blog_Topics_Database.xlsx")
    create_excel_file(topics, output_xlsx)
    print("ALL DONE!")
